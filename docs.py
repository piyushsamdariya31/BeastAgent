"""Read Word / Excel / PDF files and create Word / PDF documents."""
import os
from pathlib import Path
from xml.sax.saxutils import escape

DOC_EXT = {".pdf", ".docx", ".xlsx", ".xlsm"}


def extract(path, max_chars=60000):
    p = Path(path)
    s = p.suffix.lower()
    if s == ".pdf":
        from pypdf import PdfReader
        t = "\n\n".join((pg.extract_text() or "") for pg in PdfReader(str(p)).pages[:80])
    elif s == ".docx":
        import docx
        d = docx.Document(str(p))
        parts = [x.text for x in d.paragraphs]
        for tb in d.tables:
            parts += [" | ".join(c.text for c in row.cells) for row in tb.rows]
        t = "\n".join(parts)
    elif s in (".xlsx", ".xlsm"):
        import openpyxl
        wb, parts = openpyxl.load_workbook(str(p), read_only=True, data_only=True), []
        for ws in wb.worksheets[:6]:
            parts.append(f"## Sheet: {ws.title}")
            for r in ws.iter_rows(max_row=300, max_col=40, values_only=True):
                if any(c is not None for c in r):
                    parts.append(" | ".join("" if c is None else str(c) for c in r))
        t = "\n".join(parts)
    else:
        t = p.read_text(errors="ignore")
    return t[:max_chars]


def _rows(sec):
    t = sec.get("table") or []
    return [[str(c) for c in r] for r in t if isinstance(r, list)][:200]


def make_docx(path, title, sections):
    import docx
    d = docx.Document()
    d.add_heading(title, 0)
    for s in sections:
        if s.get("heading"):
            d.add_heading(str(s["heading"]), 1)
        for para in str(s.get("text") or "").split("\n\n"):
            if para.strip():
                d.add_paragraph(para.strip())
        for b in s.get("bullets") or []:
            d.add_paragraph(str(b), style="List Bullet")
        rows = _rows(s)
        if rows:
            tb = d.add_table(rows=len(rows), cols=max(len(r) for r in rows))
            tb.style = "Table Grid"
            for i, r in enumerate(rows):
                for j, c in enumerate(r):
                    tb.cell(i, j).text = c
    d.save(str(path))


def _font():
    from reportlab.pdfbase import pdfmetrics
    from reportlab.pdfbase.ttfonts import TTFont
    for f in ("/usr/share/fonts/truetype/dejavu/DejaVuSans.ttf", "C:/Windows/Fonts/segoeui.ttf", "C:/Windows/Fonts/arial.ttf",
              "/Library/Fonts/Arial Unicode.ttf", "/System/Library/Fonts/Supplemental/Arial Unicode.ttf"):
        if os.path.exists(f):
            pdfmetrics.registerFont(TTFont("BeastSans", f))
            return "BeastSans"
    return None


def make_pdf(path, title, sections):
    from reportlab.lib import colors
    from reportlab.lib.styles import ParagraphStyle, getSampleStyleSheet
    from reportlab.platypus import ListFlowable, Paragraph, SimpleDocTemplate, Spacer, Table, TableStyle
    font = _font()
    clean = (lambda x: str(x)) if font else (lambda x: str(x).replace("₹", "Rs.").encode("latin-1", "replace").decode("latin-1"))
    st = getSampleStyleSheet()
    if font:
        for k in ("Title", "Heading2", "BodyText"):
            st[k].fontName = font
    P = lambda t, s: Paragraph(escape(clean(t)).replace("\n", "<br/>"), st[s])
    story = [P(title, "Title"), Spacer(1, 10)]
    for s in sections:
        if s.get("heading"):
            story.append(P(s["heading"], "Heading2"))
        for para in str(s.get("text") or "").split("\n\n"):
            if para.strip():
                story += [P(para.strip(), "BodyText"), Spacer(1, 4)]
        if s.get("bullets"):
            story.append(ListFlowable([P(b, "BodyText") for b in s["bullets"]], bulletType="bullet"))
        rows = _rows(s)
        if rows:
            t = Table([[Paragraph(escape(clean(c)), st["BodyText"]) for c in r] for r in rows], repeatRows=1)
            t.setStyle(TableStyle([("GRID", (0, 0), (-1, -1), .5, colors.grey), ("BACKGROUND", (0, 0), (-1, 0), colors.lightgrey)]))
            story += [t, Spacer(1, 8)]
    SimpleDocTemplate(str(path), title=clean(title)).build(story)
